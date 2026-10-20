# -*- coding: utf-8 -*-
"""实验报告自动化 - 核心引擎
包含：SSH 免密初始化、pg_hba 本地免密、真实终端窗口采集、裁剪、Word 生成。
"""
import ctypes
import os
import re
import sys
import time
import shlex
import subprocess
import shutil
from ctypes import wintypes

# ---------------------------------------------------------------- win32 输入/窗口
K32 = ctypes.WinDLL("kernel32", use_last_error=True)
U32 = ctypes.WinDLL("user32", use_last_error=True)
try:
    ctypes.windll.shcore.SetProcessDpiAwareness(2)
except Exception:
    try:
        U32.SetProcessDPIAware()
    except Exception:
        pass


class _CHAR_UNION(ctypes.Union):
    _fields_ = [("UnicodeChar", wintypes.WCHAR), ("AsciiChar", ctypes.c_char)]


class _KEY_EVENT(ctypes.Structure):
    _fields_ = [("bKeyDown", wintypes.BOOL), ("wRepeatCount", wintypes.WORD),
                ("wVirtualKeyCode", wintypes.WORD), ("wVirtualScanCode", wintypes.WORD),
                ("uChar", _CHAR_UNION), ("dwControlKeyState", wintypes.DWORD)]


class _EVENT_UNION(ctypes.Union):
    _fields_ = [("KeyEvent", _KEY_EVENT), ("pad", ctypes.c_byte * 16)]


class _INPUT_RECORD(ctypes.Structure):
    _fields_ = [("EventType", wintypes.WORD), ("Event", _EVENT_UNION)]


def _key_records(seg):
    recs = []
    for ch in seg:
        enter = ch in "\r\n"
        for d in (1, 0):
            r = _INPUT_RECORD(); r.EventType = 1
            r.Event.KeyEvent.bKeyDown = d; r.Event.KeyEvent.wRepeatCount = 1
            r.Event.KeyEvent.wVirtualKeyCode = 0x0D if enter else 0
            r.Event.KeyEvent.uChar.UnicodeChar = "\r" if enter else ch
            recs.append(r)
    return recs


def write_input(pid, text, tries=3, chunk=480):
    """向控制台窗口写入按键（真实键盘输入）。

    控制台输入缓冲区有限（默认约 4K 事件），长文本一次性写入会失败。
    这里分块写入并校验实际写入数量，避免“命令静默丢失”。
    """
    for _ in range(tries):
        K32.FreeConsole()
        if K32.AttachConsole(pid):
            h = K32.CreateFileW("CONIN$", 0x80000000 | 0x40000000, 1 | 2, None, 3, 0, None)
            ok = True
            try:
                for i in range(0, len(text), chunk):
                    recs = _key_records(text[i:i + chunk])
                    if not recs:
                        continue
                    arr = (_INPUT_RECORD * len(recs))(*recs)
                    wr = wintypes.DWORD()
                    if (not K32.WriteConsoleInputW(h, arr, len(recs), ctypes.byref(wr))
                            or wr.value != len(recs)):
                        ok = False
                        break
                    time.sleep(0.05)
            finally:
                K32.CloseHandle(h); K32.FreeConsole()
            if ok:
                return True
        time.sleep(0.3)
    return False


_WNDENUMPROC = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)


def enum_windows():
    res = []

    @_WNDENUMPROC
    def cb(h, l):
        if U32.IsWindowVisible(h):
            ln = U32.GetWindowTextLengthW(h)
            b = ctypes.create_unicode_buffer(ln + 1); U32.GetWindowTextW(h, b, ln + 1)
            cls = ctypes.create_unicode_buffer(256); U32.GetClassNameW(h, cls, 256)
            pid = wintypes.DWORD(); U32.GetWindowThreadProcessId(h, ctypes.byref(pid))
            res.append((h, pid.value, b.value, cls.value))
        return True

    U32.EnumWindows(cb, 0)
    return res


def find_title(sub):
    for h, pid, t, c in enum_windows():
        if sub.lower() in t.lower():
            return h, pid
    return None, None


def find_pid(pid, title="LABSHOT_APP"):
    """优先按 PID 找控制台窗口；部分系统上终端窗口归属其它宿主，退回按标题匹配。"""
    fallback = None
    for h, p, t, c in enum_windows():
        if p == pid and c == "ConsoleWindowClass":
            return h
        if fallback is None and title.lower() in t.lower():
            fallback = h
    return fallback


def client_bbox(h):
    r = wintypes.RECT(); U32.GetClientRect(h, ctypes.byref(r))
    pt = wintypes.POINT(0, 0); U32.ClientToScreen(h, ctypes.byref(pt))
    return (pt.x, pt.y, pt.x + r.right, pt.y + r.bottom)


def crop_content(img, pad=12, thr=60, edge=10):
    from PIL import ImageDraw
    g = img.convert("L")
    W, H = g.size
    d = ImageDraw.Draw(g)
    d.rectangle([0, 0, W - 1, edge - 1], fill=0)
    d.rectangle([0, H - edge, W - 1, H - 1], fill=0)
    d.rectangle([0, 0, edge - 1, H - 1], fill=0)
    d.rectangle([W - edge, 0, W - 1, H - 1], fill=0)
    bb = g.point(lambda p: 255 if p > thr else 0).getbbox()
    if not bb:
        return img
    l, t, r, b = bb
    return img.crop((max(0, l - pad), max(0, t - pad),
                     min(img.width, r + pad), min(img.height, b + pad)))


_PROMPT_LINE_RE = re.compile(
    r'^\s*(?:\[[^\]]+\][#$]|[\w.-]+@[\w.-]+:[^#$]*[#$]|[A-Za-z_][\w.]*\s*=\s*#)\s*$')


# ---------------------------------------------------------------- 控制台读屏（同步）
class _COORD(ctypes.Structure):
    _fields_ = [("X", ctypes.c_short), ("Y", ctypes.c_short)]


class _SMALL_RECT(ctypes.Structure):
    _fields_ = [("Left", ctypes.c_short), ("Top", ctypes.c_short),
                ("Right", ctypes.c_short), ("Bottom", ctypes.c_short)]


class _CSBI(ctypes.Structure):
    _fields_ = [("dwSize", _COORD), ("dwCursorPosition", _COORD),
                ("wAttributes", wintypes.WORD), ("srWindow", _SMALL_RECT),
                ("dwMaximumWindowSize", _COORD)]


def read_console(pid):
    """读取目标控制台可见区域文本，用于判断命令是否执行完成。"""
    K32.FreeConsole()
    if not K32.AttachConsole(pid):
        return None
    h = K32.CreateFileW("CONOUT$", 0x80000000 | 0x40000000, 1 | 2, None, 3, 0, None)
    try:
        cs = _CSBI()
        if not K32.GetConsoleScreenBufferInfo(h, ctypes.byref(cs)):
            return None
        left, top = cs.srWindow.Left, cs.srWindow.Top
        width = cs.srWindow.Right - cs.srWindow.Left + 1
        height = cs.srWindow.Bottom - cs.srWindow.Top + 1
        n = max(width * height, 1)
        buf = ctypes.create_unicode_buffer(n + 1)
        rd = wintypes.DWORD()
        K32.ReadConsoleOutputCharacterW(h, buf, n, _COORD(left, top), ctypes.byref(rd))
        txt = buf[:rd.value]
        return "\n".join(txt[i:i + width] for i in range(0, len(txt), width))
    finally:
        try:
            K32.CloseHandle(h)
        except Exception:
            pass
        K32.FreeConsole()


def wait_settled(pid, settle=0.6, timeout=40, min_wait=0.5):
    """等待控制台画面稳定（不再变化）——比固定 sleep 更快更稳。"""
    last, stable, t0 = None, 0.0, time.time()
    while time.time() - t0 < timeout:
        cur = read_console(pid)
        if cur is not None and cur == last:
            stable += 0.15
            if stable >= settle and (time.time() - t0) >= min_wait:
                return True
        else:
            stable = 0.0
        last = cur
        time.sleep(0.15)
    return False


def grab_window(hwnd):
    """用 PrintWindow 直接抓取窗口自身像素（不受其它窗口遮挡影响）。"""
    from PIL import Image
    gdi32 = ctypes.WinDLL("gdi32")
    wr = wintypes.RECT(); U32.GetWindowRect(hwnd, ctypes.byref(wr))
    W, H = wr.right - wr.left, wr.bottom - wr.top
    if W <= 0 or H <= 0:
        return None
    hdc = U32.GetWindowDC(hwnd)
    mem = gdi32.CreateCompatibleDC(hdc)
    bmp = gdi32.CreateCompatibleBitmap(hdc, W, H)
    old = gdi32.SelectObject(mem, bmp)
    U32.PrintWindow(hwnd, mem, 2)   # PW_RENDERFULLCONTENT

    class _BMIH(ctypes.Structure):
        _fields_ = [("biSize", wintypes.DWORD), ("biWidth", wintypes.LONG), ("biHeight", wintypes.LONG),
                    ("biPlanes", wintypes.WORD), ("biBitCount", wintypes.WORD),
                    ("biCompression", wintypes.DWORD), ("biSizeImage", wintypes.DWORD),
                    ("biXPelsPerMeter", wintypes.LONG), ("biYPelsPerMeter", wintypes.LONG),
                    ("biClrUsed", wintypes.DWORD), ("biClrImportant", wintypes.DWORD)]

    class _BMI(ctypes.Structure):
        _fields_ = [("bmiHeader", _BMIH), ("bmiColors", wintypes.DWORD * 3)]

    bmi = _BMI(); bmi.bmiHeader.biSize = ctypes.sizeof(_BMIH)
    bmi.bmiHeader.biWidth = W; bmi.bmiHeader.biHeight = -H
    bmi.bmiHeader.biPlanes = 1; bmi.bmiHeader.biBitCount = 32; bmi.bmiHeader.biCompression = 0
    buf = ctypes.create_string_buffer(W * H * 4)
    gdi32.GetDIBits(mem, bmp, 0, H, buf, ctypes.byref(bmi), 0)
    img = Image.frombuffer("RGB", (W, H), buf, "raw", "BGRX", 0, 1).copy()
    gdi32.SelectObject(mem, old); gdi32.DeleteObject(bmp); gdi32.DeleteDC(mem); U32.ReleaseDC(hwnd, hdc)
    pt = wintypes.POINT(0, 0); U32.ClientToScreen(hwnd, ctypes.byref(pt))
    cr = wintypes.RECT(); U32.GetClientRect(hwnd, ctypes.byref(cr))
    ox, oy = pt.x - wr.left, pt.y - wr.top
    if img.getbbox() is None:   # 全黑说明 PrintWindow 失败
        return None
    return img.crop((ox, oy, ox + cr.right, oy + cr.bottom))


_SHELL_CMD_RE = re.compile(
    r'^\s*(?:sudo\b|su\s+-|gs_[a-z]|gsql\b|'
    r'clear\b|cd\b|ls\b|pwd\b|cat\b|echo\b|printf\b|'
    r'tar\b|gzip\b|gunzip\b|zip\b|unzip\b|bzip2\b|xz\b|'
    r'mkdir\b|rm\b|rmdir\b|cp\b|mv\b|ln\b|touch\b|chmod\b|chown\b|chgrp\b|'
    r'yum\b|dnf\b|rpm\b|apt\b|systemctl\b|service\b|'
    r'vi\b|vim\b|nano\b|sed\b|awk\b|grep\b|egrep\b|find\b|'
    r'ps\b|top\b|kill\b|free\b|df\b|du\b|iostat\b|vmstat\b|sar\b|uptime\b|'
    r'uname\b|hostname\b|ss\b|netstat\b|ifconfig\b|ip\b|ping\b|'
    r'java\b|python3?\b|wget\b|curl\b|scp\b|ssh\b|rsync\b|'
    r'export\b|source\b|env\b|which\b|whereis\b|history\b|alias\b|'
    r'useradd\b|userdel\b|passwd\b|groupadd\b|reboot\b|shutdown\b|'
    r'sysctl\b|ulimit\b|crontab\b|setenforce\b|getenforce\b)', re.I)


def is_shell_command(cmd):
    """判断步骤命令是 Shell/OS 命令（gs_checkperf、ls 等）还是 gsql 命令。"""
    c = str(cmd).strip()
    if not c or c.startswith("\\"):
        return False
    return bool(_SHELL_CMD_RE.match(c))


# ---------------------------------------------------------------- 远端（paramiko）
def _ssh(host, port, user, password=None, key=None, timeout=25):
    import paramiko
    c = paramiko.SSHClient()
    c.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    kw = dict(username=user, timeout=timeout, allow_agent=False, look_for_keys=False)
    if key:
        kw["key_filename"] = key
    if password:
        kw["password"] = password
    c.connect(host, port=port, **kw)
    return c


def remote_run(client, cmd, timeout=90):
    i, o, e = client.exec_command(cmd, timeout=timeout)
    out = o.read().decode(errors="replace")
    err = e.read().decode(errors="replace")
    return out, err


def check_environment(cfg, remote=False):
    """返回 [(项目, 是否通过, 说明)]。remote=True 时才做远程连接检测。"""
    checks = []
    ssh = shutil.which("ssh")
    checks.append(("SSH 客户端 ssh.exe", bool(ssh), ssh or "未找到（可点「环境修复」自动启用）"))
    kg = shutil.which("ssh-keygen")
    checks.append(("ssh-keygen.exe", bool(kg), kg or "未找到"))
    conhost = os.path.join(os.environ.get("SystemRoot", r"C:\Windows"), "System32", "conhost.exe")
    checks.append(("conhost.exe", os.path.exists(conhost), conhost))
    kp = cfg.get("key_path", "")
    checks.append(("SSH 私钥", bool(kp) and os.path.exists(kp), kp or "未配置（可点「① 环境修复」生成）"))
    checks.append(("Word 模板", os.path.exists(cfg.get("template", "")), cfg.get("template", "")))
    checks.append(("资料库目录", os.path.isdir(cfg.get("library", "")), cfg.get("library", "")))
    checks.append(("输出目录", os.path.isdir(cfg.get("output_dir", "")), cfg.get("output_dir", "")))
    for m in ("PIL", "docx", "paramiko"):
        try:
            __import__(m)
            checks.append(("依赖库 " + m, True, "已内置/可用"))
        except Exception as e:
            checks.append(("依赖库 " + m, False, repr(e)))
    # 远端（一次连接内完成；仅在 remote=True 时执行）
    if not remote:
        return checks
    try:
        c = _ssh(cfg.get("ssh_host", ""), int(cfg.get("ssh_port", 22)), cfg.get("ssh_user", ""),
                 password=cfg.get("ssh_password") or None,
                 key=cfg.get("key_path") if os.path.exists(cfg.get("key_path", "")) else None,
                 timeout=10)
        checks.append(("SSH 远程连通", True, cfg.get("ssh_host", "")))
        out, _ = remote_run(c, "sudo -u omm bash -lc 'source ~/.bashrc 2>/dev/null; "
                               "which gsql; echo GAUSSDATA=$GAUSSDATA'")
        gsql, data_dir = "", ""
        for line in out.splitlines():
            if line.strip().endswith("/gsql"):
                gsql = line.strip()
            elif line.startswith("GAUSSDATA="):
                data_dir = line.split("=", 1)[1].strip()
        checks.append(("远端 openGauss 已安装", bool(gsql), gsql or "未找到 gsql，可点「环境修复」"))
        sudo_out, _ = remote_run(c, "sudo -n -u omm bash -lc 'true' 2>/dev/null "
                                    "&& echo SUDO_OK || echo SUDO_NO")
        sudo_ok = "SUDO_OK" in sudo_out
        checks.append(("免密 sudo 到 omm", sudo_ok,
                       "可用" if sudo_ok else "需要密码，采集会卡在 sudo 输入处"))
        try:
            port = int(cfg.get("gsql_port") or 5432)
        except Exception:
            port = 5432
        pat = "|".join(str(x) for x in sorted({port, 5432, 26000}))
        run, _ = remote_run(c, "ss -tln 2>/dev/null | grep -E ':(%s)' | head -1" % pat)
        checks.append(("远端 openGauss 运行中", bool(run.strip()), run.strip() or "未监听，可点「环境修复」启动"))
        c.close()
    except Exception as e:
        checks.append(("SSH 远程连通", False, repr(e)))
    return checks


def ensure_ssh_key(host, port, user, password, key_path, log=print):
    """确保本地存在密钥并已安装到服务器（使用密码一次性安装）。"""
    pub = key_path + ".pub"
    keygen = shutil.which("ssh-keygen") or "ssh-keygen"
    if not os.path.exists(key_path):
        d = os.path.dirname(key_path)
        if d:
            os.makedirs(d, exist_ok=True)
        subprocess.run([keygen, "-t", "ed25519", "-N", "", "-f", key_path,
                        "-C", "labshot"], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        log("已生成本地 SSH 密钥")
    if not os.path.exists(pub):          # 只有私钥时从私钥导出公钥，避免 ssh-keygen 覆盖提示卡住
        r = subprocess.run([keygen, "-y", "-f", key_path], capture_output=True, text=True)
        if r.returncode != 0:
            raise RuntimeError("无法从私钥导出公钥：" + (r.stderr or r.stdout).strip())
        with open(pub, "w", encoding="utf-8") as f:
            f.write(r.stdout.strip() + "\n")
        log("已由私钥导出公钥")
    pubkey = open(pub, encoding="utf-8").read().strip()
    c = _ssh(host, port, user, password=password)
    try:
        remote_run(c, "mkdir -p ~/.ssh && chmod 700 ~/.ssh")
        remote_run(c, "grep -qF '%s' ~/.ssh/authorized_keys 2>/dev/null || echo '%s' >> ~/.ssh/authorized_keys"
                   % (pubkey, pubkey))
        remote_run(c, "chmod 600 ~/.ssh/authorized_keys")
        log("公钥已安装到服务器")
    finally:
        c.close()


def enable_local_trust(host, port, user, password, data_dir, log=print):
    """把 pg_hba.conf 的本地回环认证改为 trust（仅本机，带备份）。"""
    c = _ssh(host, port, user, password=password)
    try:
        remote_run(c, "sudo -u omm bash -lc 'test -f %s/pg_hba.conf.labshot || cp %s/pg_hba.conf %s/pg_hba.conf.labshot'"
                   % (data_dir, data_dir, data_dir))
        sed = (r'''sudo -u omm bash -lc 'sed -E "s#^(local[[:space:]]+all[[:space:]]+all[[:space:]]+)sha256#\1trust#; '''
               r'''s#(127\.0\.0\.1/32[[:space:]]+)sha256#\1trust#; s#(::1/128[[:space:]]+)sha256#\1trust#" '''
               r'''%s/pg_hba.conf > /tmp/pg_hba.new && cp /tmp/pg_hba.new %s/pg_hba.conf' ''' % (data_dir, data_dir))
        remote_run(c, sed)
        out, err = remote_run(c, "sudo -u omm bash -lc 'source ~/.bashrc; gs_ctl reload -D %s'" % data_dir)
        log("pg_hba 已设为本地 trust 并 reload")
    finally:
        c.close()


# ---------------------------------------------------------------- 采集
def run_capture(ssh, gsql_port, steps, workdir, setup=None, log=print, progress=None,
                db_data_dir="/opt/huawei/install/data/db", cols=120, lines=40,
                stop_db=True, capture_stop=True):
    """
    ssh: {host, port, user, key_path}
    steps: [{'name','caption','commands':[...], 'wait':sec}]
    返回: [(caption, cropped_png_path), ...]
    """
    from PIL import ImageGrab
    raw = os.path.join(workdir, "raw"); crop = os.path.join(workdir, "cropped")
    os.makedirs(raw, exist_ok=True); os.makedirs(crop, exist_ok=True)

    title = "LABSHOT_APP"
    ssh_exe = shutil.which("ssh") or "ssh"
    # cmd /c 下若首字符是引号会错误解析；无空格则不加引号，否则用 call
    ssh_cmd = ('call "%s"' % ssh_exe) if " " in ssh_exe else ssh_exe
    inner = ('title %s & mode con: cols=%d lines=%d & %s -tt -o StrictHostKeyChecking=no '
             '-o BatchMode=yes -o ConnectTimeout=20 -o ServerAliveInterval=30 '
             '-i "%s" -p %d %s@%s'
             % (title, cols, lines, ssh_cmd, ssh["key_path"], ssh["port"], ssh["user"], ssh["host"]))
    before = set(h for h, _, _, c in enum_windows() if c == "ConsoleWindowClass")
    proc = subprocess.Popen(["conhost.exe", "cmd.exe", "/c", inner],
                            creationflags=subprocess.CREATE_NEW_CONSOLE)
    h, pid = None, None
    for _ in range(80):
        for hh, pp, tt, cc in enum_windows():
            if cc == "ConsoleWindowClass" and hh not in before:
                h, pid = hh, pp
                break
        if h:
            break
        time.sleep(0.25)
    if not pid:
        raise RuntimeError("无法找到终端窗口。请把系统「默认终端应用程序」设为「Windows 控制台主机」"
                           "（conhost），不要使用 Windows Terminal 接管，然后重试。")
    U32.SetWindowPos(h, -1, 60, 20, 0, 0, 0x0001 | 0x0040)
    log("终端窗口已打开，等待登录…")

    results = []

    def shot(name):
        name = re.sub(r'[^0-9A-Za-z_.-]', '_', str(name)) or "step"
        hh = None
        for _ in range(30):
            hh = find_pid(pid)
            if hh:
                break
            time.sleep(0.2)
        if not hh:
            log("  !! 窗口丢失: " + name); return None
        for _ in range(12):                       # 若被最小化则恢复，等待有效尺寸
            if U32.IsIconic(hh):
                U32.ShowWindow(hh, 9)             # SW_RESTORE
            U32.SetWindowPos(hh, -1, 60, 20, 0, 0, 0x0001 | 0x0040)
            time.sleep(0.3)
            r = wintypes.RECT(); U32.GetClientRect(hh, ctypes.byref(r))
            if r.right > 20 and r.bottom > 20:
                break
        im = grab_window(hh)
        if im is None or im.getbbox() is None:
            im = ImageGrab.grab(bbox=client_bbox(hh))
        p1 = os.path.join(raw, name + ".png"); im.save(p1)
        p2 = os.path.join(crop, name + ".png"); crop_content(im).save(p2)
        return p2

    def send(text):
        if not write_input(pid, text):
            raise RuntimeError("向终端写入命令失败（窗口可能已关闭）：" + " ".join(text.split())[:80])

    def screen():
        return read_console(pid) or ""

    def login_error(scr):
        for k in ("Permission denied", "Connection refused", "Connection timed out",
                  "Could not resolve", "Host key verification failed", "No route to host",
                  "password:", "Password:", "[sudo] password for",
                  "not in the sudoers file", "is not allowed to execute"):
            if k in scr:
                return k
        return ""

    def wait_prompt(timeout=35, min_wait=0.6):
        """等到屏幕最后一行出现 shell 提示符，避免把命令打进密码提示/连接中。"""
        t0, scr = time.time(), ""
        while time.time() - t0 < timeout:
            scr = screen()
            lines = [x.rstrip() for x in scr.splitlines() if x.strip()]
            if (time.time() - t0) >= min_wait and lines and re.search(r'[#$>]\s*$', lines[-1]):
                return True, scr
            time.sleep(0.2)
        return False, scr

    def wait_omm_prompt(timeout=20):
        """确认已经切换到 omm 用户（提示符以 [omm@ 开头）。"""
        t0, scr = time.time(), ""
        while time.time() - t0 < timeout:
            scr = screen()
            lines = [x.rstrip() for x in scr.splitlines() if x.strip()]
            if lines and re.match(r'^\s*\[omm@[^\]]*\][#$]\s*$', lines[-1]):
                return True, scr
            if login_error(scr):
                return False, scr
            time.sleep(0.2)
        return False, scr

    def wait_gsql_prompt(timeout=45):
        """等到 gsql 提示符（openGauss=# / db1=#），而不是 shell 提示符。"""
        t0, scr = time.time(), ""
        while time.time() - t0 < timeout:
            scr = screen()
            lines = [x.rstrip() for x in scr.splitlines() if x.strip()]
            if lines and re.match(r'^[A-Za-z_][\w.]*\s*=\s*#', lines[-1]):
                return True, scr
            if any(k in scr for k in ("failed to connect", "FATAL",
                                      "command not found", "No such file")):
                return False, scr
            time.sleep(0.2)
        return False, scr

    def wait_done(timeout=180, min_wait=0.6):
        """等到最后一行重新出现提示符（shell 的 $/# 或 gsql 的 =#），才认为命令执行完。"""
        t0 = time.time()
        while time.time() - t0 < timeout:
            scr = screen()
            lines = [x.rstrip() for x in scr.splitlines() if x.strip()]
            if (time.time() - t0) >= min_wait and lines and _PROMPT_LINE_RE.match(lines[-1]):
                return True
            time.sleep(0.15)
        return False

    def send_wait(text, timeout=180, min_wait=0.6):
        send(text)
        if not wait_done(timeout=timeout, min_wait=min_wait):
            raise RuntimeError("命令执行超时/未返回提示符（%s）。屏幕内容：\n%s"
                               % (" ".join(text.split())[:80], screen()[-400:]))

    def enter_gsql(timeout=45):
        send("gsql -d postgres -p %d -r\r" % gsql_port)
        ok, scr = wait_gsql_prompt(timeout)
        tail = scr.rsplit("gsql -d postgres", 1)[-1] if "gsql -d postgres" in scr else scr
        if not ok:
            raise RuntimeError("打开数据库失败（请检查数据库是否启动/端口是否正确）。屏幕内容：\n"
                               + tail[-400:])
        send_wait("\\pset pager off\r", timeout=20)   # 关闭分页器，避免 --More-- 卡住

    try:
        ok, scr = wait_prompt(35)
        if not ok:
            raise RuntimeError("SSH 登录失败/超时（%s）。请先点「① 环境修复与初始化」。屏幕内容：\n%s"
                               % (login_error(scr) or "未出现提示符", scr[-400:]))
        send("sudo -u omm bash -l\r")
        ok, scr = wait_omm_prompt(20)
        if not ok:
            raise RuntimeError("切换到 omm 用户失败（%s）。屏幕内容：\n%s"
                               % (login_error(scr) or "未出现提示符", scr[-400:]))
        # 优先 gs_om（gs_checkperf 依赖 Normal 集群状态），失败再回退 gs_ctl
        send_wait("gs_om -t start || gs_ctl start -D %s -Z single_node -l %s/logfile\r"
                  % (db_data_dir, db_data_dir), timeout=180, min_wait=1.0)
        send_wait("clear\r", timeout=15, min_wait=0.4)   # 清掉启动输出
        enter_gsql()

        for c in (setup or []):      # 连接后、截图前的清理/准备（不截图）
            send_wait(c + "\r", timeout=120)

        in_sql = True
        total = len(steps)
        for i, st in enumerate(steps):
            send_wait("\\! clear\r" if in_sql else "clear\r", timeout=15, min_wait=0.4)
            for c in st.get("commands", []):
                if st.get("shell") or is_shell_command(c):
                    if in_sql:                      # 退出 gsql，Shell 命令要在 omm shell 里跑
                        send_wait("\\q\r", timeout=30)
                        in_sql = False
                    send_wait(c + "\r", timeout=180)   # gs_checkperf 等可能跑 15~30 秒
                else:
                    if not in_sql:                  # 重新连接 gsql 再执行 SQL
                        enter_gsql()
                        in_sql = True
                    send_wait(c + "\r", timeout=90)
            p = shot(st.get("name", "step%02d" % (i + 1)))
            if p:
                results.append((st.get("caption", ""), p))
            if progress:
                progress(int((i + 1) / total * 100))
            log("  [%d/%d] %s" % (i + 1, total, st.get("caption", "")))

        if stop_db:
            send_wait("\\! clear\r" if in_sql else "clear\r", timeout=15, min_wait=0.4)
            if in_sql:
                send_wait("\\q\r", timeout=30)
                in_sql = False
            send_wait("gs_om -t stop || gs_ctl stop -D %s\r" % db_data_dir,
                      timeout=180, min_wait=1.0)
            if capture_stop:
                p = shot("99_quit_stop")
                if p:
                    results.append(("退出数据库连接并停止 openGauss 服务", p))
            log("  已退出并停止数据库")
    finally:
        try:
            hh = find_pid(pid)
            if hh:
                U32.PostMessageW(hh, 0x0010, 0, 0)   # WM_CLOSE 关闭终端窗口
        except Exception:
            pass
        try:
            proc.terminate()
        except Exception:
            pass
        try:
            proc.kill()
        except Exception:
            pass
    return results


# ---------------------------------------------------------------- Word 生成
def build_docx(template, out_path, exp_name, purpose, device, result, problem, feeling,
               image_items, cover=None, log=print):
    """image_items: [(caption, png_path), ...]；cover: 封面字段 dict"""
    import docx
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from PIL import Image

    def set_font(run, name="仿宋", size=11, bold=False):
        run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
        run.font.underline = False
        rPr = run._element.get_or_add_rPr()
        rf = rPr.find(qn('w:rFonts'))
        if rf is None:
            rf = rPr.makeelement(qn('w:rFonts'), {}); rPr.append(rf)
        for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            rf.set(qn(a), name)

    HINTS = ["（题目格式", "每进行1次实训", "注释：学生总结", "格式：黑体", "（仿宋小四）", "（标题为宋体"]
    doc = docx.Document(template)
    for p in list(doc.paragraphs):
        if any(h in p.text for h in HINTS):
            p._element.getparent().remove(p._element)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if any(h in p.text for h in HINTS):
                        for r in list(p.runs):
                            r._element.getparent().remove(r._element)

    report_tbl = None
    for t in doc.tables:
        if len(t.columns) == 2 and any("过程" in c.text for r in t.rows for c in r.cells):
            report_tbl = t; break
    if report_tbl is None:
        raise RuntimeError("模板中未找到报告表")

    # 填写封面（课程名称/姓名/学号/专业/班级/时间/地点）
    if cover:
        norm = lambda s: s.replace(" ", "").replace("\u3000", "").strip()
        for tb in doc.tables:
            if len(tb.columns) >= 3:
                keys = [norm(tb.rows[i].cells[0].text) for i in range(len(tb.rows))]
                if any("课程名称" in k for k in keys):
                    for i in range(len(tb.rows)):
                        key = norm(tb.rows[i].cells[0].text)
                        if key in cover:
                            cell = tb.rows[i].cells[2]
                            for p in list(cell.paragraphs):
                                p._element.getparent().remove(p._element)
                            cell.add_paragraph()
                            p = cell.paragraphs[0]
                            for j, line in enumerate(str(cover[key]).split("\n")):
                                if j:
                                    p.add_run().add_break()
                                set_font(p.add_run(line))
                    break

    report_tbl.rows[0].cells[1].paragraphs[0].text = ""
    set_font(report_tbl.rows[0].cells[1].paragraphs[0].add_run(
        time.strftime("%Y 年 %m 月 %d 日")))
    nc = report_tbl.rows[1].cells[1]
    for p in list(nc.paragraphs):
        p._element.getparent().remove(p._element)
    nc.add_paragraph()
    nc.paragraphs[0].text = ""
    set_font(nc.paragraphs[0].add_run(exp_name))

    for p in doc.paragraphs:
        if "实训（验）报告（" in p.text:
            p.text = ""
            set_font(p.add_run("实训（验）报告（%s）" % exp_name), size=18, bold=True)

    content = report_tbl.rows[2].cells[1]
    for p in list(content.paragraphs):
        p._element.getparent().remove(p._element)
    content.add_paragraph()

    def heading(txt):
        p = content.add_paragraph(); p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(txt), name="宋体", size=12, bold=True)

    def body(txt):
        p = content.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.4
        set_font(p.add_run(txt), name="仿宋", size=11)

    heading("1.1 实训（验）目的与要求"); body(purpose)
    heading("1.2 实训（验）设备与方法"); body(device)
    heading("1.3 实训内容与步骤")
    body("本部分依次完成下列操作，具体命令与运行结果如下图所示。")
    for cap, png in image_items:
        p = content.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        set_font(p.add_run(cap), size=10.5, bold=True)
        with Image.open(png) as im:
            w, hh = im.size
        ip = content.add_paragraph(); ip.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ip.paragraph_format.space_after = Pt(10)
        ip.add_run().add_picture(png, width=Cm(min(w / 96.0 * 2.54, 16.0)))
    heading("1.4 实训结果与分析"); body(result)
    heading("1.5 遇到的问题与解决方法"); body(problem)
    heading("1.6 实训心得与体会"); body(feeling)

    doc.save(out_path)
    log("已生成 Word：" + out_path)
    return out_path

# ---------------------------------------------------------------- 密码加密（DPAPI, D）
class _BLOB(ctypes.Structure):
    _fields_ = [("cbData", wintypes.DWORD), ("pbData", ctypes.POINTER(ctypes.c_char))]


def _mkblob(data):
    buf = ctypes.create_string_buffer(data, len(data) + 1)
    return _BLOB(len(data), ctypes.cast(buf, ctypes.POINTER(ctypes.c_char))), buf


def dpapi_encrypt(text):
    import base64
    if not text:
        return ""
    data = text.encode("utf-8")
    bin_, hold = _mkblob(data)
    out = _BLOB()
    if not ctypes.windll.crypt32.CryptProtectData(ctypes.byref(bin_), None, None, None, None, 0, ctypes.byref(out)):
        raise ctypes.WinError()
    try:
        return "enc:" + base64.b64encode(ctypes.string_at(out.pbData, out.cbData)).decode()
    finally:
        ctypes.windll.kernel32.LocalFree(out.pbData)


def dpapi_decrypt(token):
    import base64
    if not token:
        return ""
    if not token.startswith("enc:"):
        return token  # 兼容明文
    data = base64.b64decode(token[4:])
    bin_, hold = _mkblob(data)
    out = _BLOB()
    if not ctypes.windll.crypt32.CryptUnprotectData(ctypes.byref(bin_), None, None, None, None, 0, ctypes.byref(out)):
        raise ctypes.WinError()
    try:
        return ctypes.string_at(out.pbData, out.cbData).decode("utf-8")
    finally:
        ctypes.windll.kernel32.LocalFree(out.pbData)


# ---------------------------------------------------------------- 远程环境探测（C）
def detect_remote_env(ssh, password=None, log=print):
    """探测远端 openGauss：gsql 路径、GAUSSHOME/GAUSSDATA、端口。"""
    env = {"gsql": "", "gauss_home": "", "data_dir": "", "port": "", "installed": False}
    try:
        c = _ssh(ssh["host"], ssh["port"], ssh["user"], password=password,
                 key=ssh.get("key_path") if os.path.exists(ssh.get("key_path", "")) else None, timeout=12)
    except Exception as e:
        log("远程探测：连接失败 " + repr(e))
        return env
    try:
        out, _ = remote_run(c, "sudo -u omm bash -lc 'source ~/.bashrc 2>/dev/null; "
                               "echo HOME=$HOME; which gsql; echo GAUSSHOME=$GAUSSHOME; "
                               "echo GAUSSDATA=$GAUSSDATA; gsql --version 2>/dev/null | head -1'")
        for line in out.splitlines():
            line = line.strip()
            if line.endswith("bin/gsql") or line.endswith("/gsql"):
                env["gsql"] = line
            elif line.startswith("GAUSSHOME="):
                env["gauss_home"] = line.split("=", 1)[1]
            elif line.startswith("GAUSSDATA="):
                env["data_dir"] = line.split("=", 1)[1]
        env["installed"] = bool(env["gsql"])
        if not env["data_dir"]:          # GAUSSDATA 未设置时，从进程参数里取 -D 数据目录
            po, _ = remote_run(c, "ps -eo args 2>/dev/null | grep '[g]aussdb' | head -1")
            m = re.search(r'-D\s+(\S+)', po)
            if m:
                env["data_dir"] = m.group(1)
        if env["data_dir"]:
            po, _ = remote_run(c, "sudo -u omm bash -lc 'grep -E \"^[[:space:]]*port[[:space:]]*=\" %s/postgresql.conf 2>/dev/null | head -1'"
                               % env["data_dir"])
            m = re.search(r"port\s*=\s*(\d+)", po)
            if m:
                env["port"] = m.group(1)
        if not env["port"]:
            so, _ = remote_run(c, "ss -tln 2>/dev/null | grep -oE ':(5432|26000|15400)' | head -1")
            m = re.search(r":(\d+)", so)
            if m:
                env["port"] = m.group(1)
        log("远程环境：gsql=%s, GAUSSDATA=%s, port=%s, installed=%s"
            % (env["gsql"] or "-", env["data_dir"] or "-", env["port"] or "-", env["installed"]))
    except Exception as e:
        log("远程探测异常：" + repr(e))
    finally:
        c.close()
    return env


def ensure_remote_running(ssh, password=None, data_dir="", log=print, port=5432):
    """确保 openGauss 以 gs_om 正常集群状态运行（gs_checkperf 等工具依赖此状态）。"""
    try:
        port = int(port or 5432)
    except Exception:
        port = 5432
    c = _ssh(ssh["host"], ssh["port"], ssh["user"], password=password,
             key=ssh.get("key_path") if os.path.exists(ssh.get("key_path", "")) else None)

    def run_omm(cmd, timeout=180):
        return remote_run(c, "sudo -u omm bash -lc %s"
                          % shlex.quote("source ~/.bashrc 2>/dev/null; " + cmd), timeout=timeout)

    def listening():
        out, _ = remote_run(c, "ss -tln 2>/dev/null | grep -E ':%d\\b' | head -1" % port)
        return bool(out.strip())

    def om_normal():
        out, _ = run_omm("gs_om -t status --detail 2>&1")
        return bool(re.search(r'cluster_state\s*:\s*Normal', out)), out

    try:
        normal, st = om_normal()
        if listening() and normal:
            log("数据库已在运行（端口 %d，集群状态 Normal）" % port)
            return True
        log("数据库未处于正常集群状态，尝试用 gs_om 启动…")
        o2, e2 = run_omm("gs_om -t start 2>&1")
        time.sleep(3)
        normal, st = om_normal()
        if listening() and normal:
            log("启动成功（端口 %d，集群状态 Normal）" % port)
            return True
        if not data_dir:
            log("gs_om 启动失败，且未配置数据目录，无法回退 gs_ctl：" + (e2.strip() or o2.strip())[:300])
            return False
        log("gs_om 启动未成功，回退 gs_ctl…")
        o3, e3 = run_omm("gs_ctl start -D %s -Z single_node -l %s/logfile" % (data_dir, data_dir),
                         timeout=120)
        time.sleep(3)
        ok = listening()
        log("启动" + ("成功（端口 %d）" % port if ok else
                      "失败：" + (e3.strip() or o3.strip() or e2.strip() or o2.strip())[:300]))
        return ok
    finally:
        c.close()


def reset_databases(ssh, password=None, port=5432, statements=None, log=print):
    """数据库初始化：在 postgres 库中逐条执行 setup（DROP 库/用户等）。

    每条语句单独执行，避免多语句事务导致 DROP DATABASE 失败；
    返回是否全部成功。
    """
    items, seen = [], set()
    for raw in statements or []:
        for part in str(raw).split(";"):        # setup 均为 DROP 语句，按分号拆开逐条执行
            s = " ".join(part.split()).strip()
            if not s:
                continue
            if not s.endswith(";"):
                s += ";"
            if s not in seen:
                seen.add(s); items.append(s)
    if not items:
        log("数据库初始化：所选实验没有可执行的 setup 语句")
        return True
    try:
        port = int(port or 5432)
    except Exception:
        port = 5432
    c = _ssh(ssh["host"], ssh["port"], ssh["user"], password=password,
             key=ssh.get("key_path") if os.path.exists(ssh.get("key_path", "")) else None)
    ok = True
    try:
        for s in items:
            inner = ("source ~/.bashrc 2>/dev/null; "
                     "gsql -d postgres -p %d -r -c %s < /dev/null" % (port, shlex.quote(s)))
            cmd = "sudo -u omm bash -lc %s" % shlex.quote(inner)
            out, err = remote_run(c, cmd, timeout=120)
            text = " ".join((out + "\n" + err).split())
            if "ERROR" in text.upper():
                ok = False
                log("  [×] %s -> %s" % (s, text[:240]))
            else:
                log("  [√] %s" % s)
    finally:
        c.close()
    return ok


# ---------------------------------------------------------------- 本机 OpenSSH 修复（A）
def enable_windows_openssh(log=print):
    """尝试启用 Windows OpenSSH 客户端（需管理员）。返回是否成功。"""
    if shutil.which("ssh") and shutil.which("ssh-keygen"):
        log("本机已有 OpenSSH 客户端");
        return True
    log("未检测到 OpenSSH 客户端，尝试启用（需要管理员权限）…")
    ps = ("Add-WindowsCapability -Online -Name OpenSSH.Client~~~~0.0.1.0")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                           capture_output=True, text=True, timeout=180)
        out = (r.stdout or "") + (r.stderr or "")
        log(out.strip()[-500:])
        ok = bool(shutil.which("ssh"))
        if not ok:
            log("启用失败。可手动：设置→应用→可选功能→添加“OpenSSH 客户端”，或以管理员运行上面的命令。")
        return ok
    except Exception as e:
        log("启用失败：" + repr(e))
        log("可手动下载：https://github.com/PowerShell/Win32-OpenSSH/releases")
        return False
